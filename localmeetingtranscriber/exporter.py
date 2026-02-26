"""DOCX export utilities."""
from __future__ import annotations

import re
from pathlib import Path
from docx import Document


def _sanitize_title(title: str) -> str:
    cleaned = " ".join(title.strip().split())
    cleaned = re.sub(r"[^\w\-\s\.\(\)\u4e00-\u9fff]", "", cleaned)
    return cleaned or "Meeting"


def build_output_path(output_dir: Path, title: str) -> Path:
    """Build a safe output path for the DOCX file."""
    safe_title = _sanitize_title(title)
    return output_dir / f"{safe_title}.docx"


def export_docx(
    output_path: Path,
    title: str,
    date: str,
    lines: list[str],
    appendix: list[str] | None = None,
) -> None:
    """Export transcript content to a DOCX file."""
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Date: {date}")
    doc.add_paragraph("")
    for line in lines:
        doc.add_paragraph(line)

    if appendix:
        doc.add_page_break()
        doc.add_heading("Appendix: Raw Transcript", level=2)
        for line in appendix:
            doc.add_paragraph(line)

    doc.save(output_path)
